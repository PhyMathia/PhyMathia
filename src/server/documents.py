"""文件上传与文档/图片知识点提取。"""

import io
import json
import logging
import os
import re
import time
from pathlib import Path

from http_client import get_http_client

from .config import AI_PROVIDERS, LEVEL_PROMPTS, OPENCODE_DEFAULT_API_KEY, UPLOAD_DIR, UPLOADS_META_PATH
from .knowledge import _clean_knowledge_title, _looks_like_formula, _normalize_formula
from .storage import _read_json, _write_json

logger = logging.getLogger(__name__)

def _sanitize_filename(filename: str) -> str:
    name = Path(str(filename or "upload")).name
    name = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()[:120] or "upload"
    return name


def _read_upload(file_id: str):
    meta = _read_json(UPLOADS_META_PATH, {})
    entry = meta.get(file_id)
    if not entry:
        return None
    path = UPLOAD_DIR / f"{file_id}.bin"
    if not path.exists():
        return None
    return entry, path.read_bytes()


def _save_upload(file_id: str, filename: str, content: bytes) -> None:
    meta = _read_json(UPLOADS_META_PATH, {})
    meta[file_id] = {
        "id": file_id,
        "filename": filename,
        "savedAt": int(time.time() * 1000),
    }
    _write_json(UPLOADS_META_PATH, meta)
    (UPLOAD_DIR / f"{file_id}.bin").write_bytes(content)


def _extract_image_text(content: bytes) -> str:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        return ""
    try:
        tmp_path = UPLOAD_DIR / "ocr_tmp.png"
        tmp_path.write_bytes(content)
        result, _ = RapidOCR()(str(tmp_path))
        if result:
            return "\n".join(str(item[1]) for item in result if len(item) > 1)
    except Exception as e:
        logger.warning(f"Image OCR failed: {e}")
    return ""


def _extract_document_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".tex"}:
        for enc in ("utf-8", "utf-8-sig", "gbk"):
            try:
                return content.decode(enc)
            except (UnicodeDecodeError, ValueError):
                continue
        return content.decode("utf-8", "ignore")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return ""
        try:
            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:
            logger.warning(f"PDF extraction failed: {e}")
            return ""
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            return ""
        try:
            doc = docx.Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            logger.warning(f"DOCX extraction failed: {e}")
            return ""
    if ext in {".pptx", ".ppt"}:
        try:
            from pptx import Presentation
        except ImportError:
            return ""
        try:
            prs = Presentation(io.BytesIO(content))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
                    if hasattr(shape, "has_table") and shape.has_table:
                        for row in shape.table.rows:
                            parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            logger.warning(f"PPTX extraction failed: {e}")
            return ""
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff"}:
        return _extract_image_text(content)
    return ""


def _parse_document_extract_json(text: str, max_items: int):
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", text or "")
    if m:
        text = m.group(1)
    else:
        start, end = text.find("{"), text.rfind("}")
        if start >= 0 and end > start:
            text = text[start:end + 1]
    try:
        data = json.loads(text)
    except (json.JSONDecodeError, TypeError):
        return [], [], []
    if not isinstance(data, dict):
        return [], [], []
    nodes = data.get("nodes") or data.get("items") or []
    edges = data.get("edges") or []
    raw_relations = data.get("relations") or []
    normalized = []
    for it in nodes if isinstance(nodes, list) else []:
        if not isinstance(it, dict):
            continue
        title = _clean_knowledge_title(str(it.get("title") or ""))
        if not title:
            continue
        category = it.get("category")
        if category not in ("physics", "math", "other"):
            category = "other"
        formulas = []
        for f in (it.get("formulas") or []):
            fs = _normalize_formula(str(f).strip())
            if fs and _looks_like_formula(fs):
                formulas.append(fs)
        normalized.append({
            "id": str(it.get("id") or it.get("key") or title),
            "title": title[:80],
            "category": category,
            "tags": [str(t).strip() for t in (it.get("tags") or []) if str(t).strip()][:6],
            "summary": str(it.get("summary") or "")[:240],
            "formulas": formulas[:8],
        })
        for rel in (it.get("relations") or []):
            if not isinstance(rel, dict):
                continue
            edges.append({
                "from": str(it.get("id") or it.get("key") or title),
                "to": str(rel.get("target") or rel.get("to") or ""),
                "type": str(rel.get("type") or "related"),
                "label": str(rel.get("label") or rel.get("type") or "相关"),
                "score": rel.get("score"),
            })

    by_key = {}
    result_nodes = []
    for index, item in enumerate(normalized[:max_items]):
        raw_id = str(item["id"])
        node_id = f"n{index + 1}"
        by_key[raw_id] = node_id
        result_nodes.append({**item, "id": node_id})

    result_edges = []
    for edge in edges if isinstance(edges, list) else []:
        if not isinstance(edge, dict):
            continue
        from_key = str(edge.get("from") or edge.get("source") or "")
        to_key = str(edge.get("to") or edge.get("target") or "")
        from_id = by_key.get(from_key)
        to_id = by_key.get(to_key)
        if not from_id or not to_id or from_id == to_id:
            continue
        edge_type = str(edge.get("type") or "related")[:40]
        result_edges.append({
            "from": from_id,
            "to": to_id,
            "type": edge_type,
            "label": str(edge.get("label") or edge_type)[:120],
            "score": _safe_score(edge.get("score")),
        })

    result_relations = []
    for rel in raw_relations if isinstance(raw_relations, list) else []:
        if not isinstance(rel, dict):
            continue
        node_keys = rel.get("nodes") or rel.get("nodeIds") or []
        if not node_keys and rel.get("from") and rel.get("to"):
            node_keys = [rel.get("from"), rel.get("to")]
        mapped = []
        for key in node_keys if isinstance(node_keys, list) else []:
            mapped_id = by_key.get(str(key))
            if mapped_id and mapped_id not in mapped:
                mapped.append(mapped_id)
        if len(mapped) < 2:
            continue
        label = str(rel.get("label") or rel.get("meaning") or rel.get("type") or "")
        result_relations.append({
            "nodes": mapped[:4],
            "label": label[:200],
            "type": str(rel.get("type") or "本质联系")[:40],
            "score": _safe_score(rel.get("score"), 0.75),
        })
    result_relations.sort(key=lambda item: item["score"], reverse=True)
    return result_nodes, result_edges[:max_items * 4], result_relations[:3]


def _safe_score(value, default: float = 0.7) -> float:
    try:
        score = float(value)
        if 0 <= score <= 1:
            return score
    except (TypeError, ValueError):
        pass
    return default


def _extract_text_formulas(text: str) -> list:
    formulas = []
    if not text:
        return formulas
    for m in re.finditer(r"\$\$([^$\n]+)\$\$|\\\((.+?)\\\)|\\\[(.+?)\\\]|\$([^$\n]+)\$", text):
        expr = next((g for g in m.groups() if g), "")
        normalized = _normalize_formula(expr)
        if normalized and _looks_like_formula(normalized) and normalized not in formulas:
            formulas.append(normalized)
        if len(formulas) >= 8:
            break
    return formulas


def _local_extract_document_knowledge(text: str, filename: str, max_items: int):
    if not text or not text.strip():
        return [], [], []
    headings = list(re.finditer(r"^#{1,6}\s+(.+?)\s*$", text, re.M))
    if headings:
        nodes = []
        stack = []
        for idx, match in enumerate(headings):
            title = _clean_knowledge_title(match.group(1))
            level = len(match.group(0)) - len(match.group(0).lstrip("#"))
            start = match.end()
            end = headings[idx + 1].start() if idx + 1 < len(headings) else len(text)
            section = text[start:end]
            summary = re.sub(r"\s+", " ", section).strip()[:200]
            if not summary:
                summary = title
            category = "other"
            if any(k in section for k in ("物理", "力学", "电磁", "光学", "热", "振动", "波", "场", "力", "能量")):
                category = "physics"
            elif any(k in section for k in ("方程", "函数", "导数", "积分", "矩阵", "几何", "代数", "微分", "定理", "证明", "数学")):
                category = "math"
            node = {
                "id": f"local_{idx + 1}",
                "title": title[:80],
                "category": category,
                "tags": ["物理" if category == "physics" else "数学" if category == "math" else "其他"],
                "summary": summary,
                "formulas": _extract_text_formulas(section),
            }
            nodes.append(node)
            while stack and stack[-1]["level"] >= level:
                stack.pop()
            if stack:
                edge_label = "包含"
                if not any(e["from"] == stack[-1]["id"] and e["to"] == node["id"] for e in nodes[-1].get("_edges", [])):
                    node.setdefault("_edges", []).append({"from": stack[-1]["id"], "to": node["id"], "type": "包含", "label": edge_label})
            stack.append({"id": node["id"], "level": level})
        edges = []
        for node in nodes:
            edges.extend(node.get("_edges", []))
            node.pop("_edges", None)
        return nodes[:max_items], edges[:max_items * 4], []

    paragraphs = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", text) if len(p.strip()) >= 10]
    if not paragraphs:
        paragraphs = [re.sub(r"\s+", " ", text).strip()[:1000]]
    nodes = []
    for index, paragraph in enumerate(paragraphs[:max_items]):
        title = paragraph[:40] or f"{Path(filename).stem} 知识点 {index + 1}"
        category = "other"
        if any(k in paragraph for k in ("物理", "力学", "电磁", "光学", "热", "振动", "波", "场", "力", "能量")):
            category = "physics"
        elif any(k in paragraph for k in ("方程", "函数", "导数", "积分", "矩阵", "几何", "代数", "微分", "定理", "证明", "数学")):
            category = "math"
        nodes.append({
            "id": f"para_{index + 1}",
            "title": title[:80],
            "category": category,
            "tags": ["物理" if category == "physics" else "数学" if category == "math" else "其他"],
            "summary": paragraph[:200],
            "formulas": _extract_text_formulas(paragraph),
        })
    return nodes, [], []


async def _ai_extract_document_knowledge(
    text: str,
    filename: str,
    is_image: bool,
    image_b64: str,
    provider: str,
    api_key: str,
    model: str,
    base_url: str,
    level: str,
    max_items: int,
):
    if not base_url:
        base_url = AI_PROVIDERS.get(provider, {}).get("base_url", "")
    if not base_url:
        return [], [], []
    if not api_key and provider == "opencode-go":
        api_key = os.getenv("OPENCODE_GO_API_KEY", "") or os.getenv("OPENCODE_API_KEY", "")
    if not api_key and provider == "opencode":
        api_key = OPENCODE_DEFAULT_API_KEY
    if not api_key and provider not in ("opencode", "opencode-go"):
        return [], [], []

    prompt = (
        "你是 PhyMathia 的文件知识抽取助手。请从用户上传的文件中抽取最重要的知识点。\n"
        f"要求：最多输出 {max_items} 个知识点；每个知识点给出具体名称、分类、简短摘要和公式；"
        "不要抽取知识点之间的关系（关系由后续 AI 网络助手按用户意图整理）。\n"
        "抽取时优先覆盖不同主题/段落：著名的定律、定理、公式（如麦克斯韦方程组、牛顿定律、欧姆定律等）不应遗漏；如果材料含多个主题，每个主题至少抽取 1 条，不要集中在前半部分。\n"
        '只输出 JSON，不要输出其他内容：\n'
        '{"nodes":[{"id":"n1","title":"知识点名称","category":"physics|math|other","tags":["标签"],"summary":"一句话摘要","formulas":["$F=ma$"]}]}'
    )
    level_suffix = LEVEL_PROMPTS.get(level, LEVEL_PROMPTS["university"])
    user_content = f"文件名：{filename}\n\n{text[:12000]}" if text else f"文件名：{filename}"
    if is_image and image_b64 and not text:
        ext = Path(filename).suffix.lower()
        mime = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
            ".tiff": "image/tiff",
        }.get(ext, "image/png")
        message_content = [
            {"type": "text", "text": user_content},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{image_b64}"}},
        ]
    else:
        message_content = user_content
    messages = [
        {"role": "system", "content": prompt + "\n\n难度要求：" + level_suffix},
        {"role": "user", "content": message_content},
    ]
    url = f"{base_url.rstrip('/')}/chat/completions"
    headers = {"Content-Type": "application/json"}
    if provider != "opencode":
        headers["Authorization"] = f"Bearer {api_key}"
    body = {"model": model, "messages": messages, "stream": False, "temperature": 0.2}
    client = get_http_client()
    resp = await client.post(url, json=body, headers=headers, timeout=120.0)
    resp.raise_for_status()
    data = resp.json()
    content = data["choices"][0]["message"]["content"]
    return _parse_document_extract_json(content, max_items)




__all__ = [
    "_sanitize_filename", "_read_upload", "_save_upload", "_extract_document_text",
    "_ai_extract_document_knowledge", "_local_extract_document_knowledge",
]
